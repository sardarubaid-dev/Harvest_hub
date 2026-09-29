import os
import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--break-system-packages"])
    import docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

PRIMARY_GREEN = RGBColor(0x1E, 0x3C, 0x2F)
DARK_GREEN = RGBColor(0x2E, 0x7D, 0x32)
ORANGE = RGBColor(0xF5, 0x7C, 0x00)
DARK_TEXT = RGBColor(0x21, 0x21, 0x21)
GREY_TEXT = RGBColor(0x61, 0x61, 0x61)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shading_elm = docx.oxml.OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), color_hex)
    shading_elm.set(qn('w:val'), 'clear')
    shading.append(shading_elm)

def styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = PRIMARY_GREEN
    return h

def styled_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        set_cell_shading(cell, '1E3C2F')
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True
                run.font.size = Pt(10)
    for row_data in rows:
        row = table.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = val
            for p in row.cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    return table

def bold_line(doc, label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.color.rgb = DARK_TEXT
    r2 = p.add_run(value)
    r2.font.color.rgb = GREY_TEXT

def bullet(doc, text):
    p = doc.add_paragraph(text, style='List Bullet')
    for run in p.runs:
        run.font.size = Pt(11)

def create_doc():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.font.color.rgb = DARK_TEXT
    style.paragraph_format.space_after = Pt(2)
    style.paragraph_format.space_before = Pt(2)

    for i in range(1, 4):
        hs = doc.styles[f'Heading {i}']
        hs.paragraph_format.space_before = Pt(8)
        hs.paragraph_format.space_after = Pt(4)

    # ==================== COVER PAGE ====================
    for _ in range(6):
        doc.add_paragraph()
    
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run('HarvestHub')
    r.font.size = Pt(42)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_GREEN

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('Marketplace for Local Farm Products')
    r.font.size = Pt(18)
    r.font.color.rgb = ORANGE

    
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = line.add_run('Comprehensive Project Report & Documentation')
    r.font.size = Pt(14)
    r.font.color.rgb = GREY_TEXT


    info_items = [
        'Course: Multi-Platform App Computing',
        'Platform: Cross-Platform (Android & iOS)',
        'Framework: Flutter SDK (Dart)',
        'Backend: Firebase (Firestore, Auth, Storage)',
    ]
    for item in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(item)
        r.font.size = Pt(12)
        r.font.color.rgb = DARK_TEXT

    img_path = '/home/muhammad-hanzala/.gemini/antigravity/brain/857111e4-9b2e-402d-98a9-e323d6d0d14a/.user_uploaded/media_1790553653791.jpg'
    if os.path.exists(img_path):
        doc.add_paragraph()
        last_p = doc.add_paragraph()
        last_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = last_p.add_run()
        run.add_picture(img_path, width=Inches(3.0))

    doc.add_page_break()

    # ==================== TABLE OF CONTENTS ====================
    styled_heading(doc, 'Table of Contents', 1)
    toc_items = [
        '1. Problem Definition',
        '2. Design Specifications',
        '   2.1 Technology Stack',
        '   2.2 System Architecture',
        '   2.3 Color Theme & UI/UX Guidelines',
        '3. Diagrams',
        '   3.1 Application Flow Diagram',
        '   3.2 User Authentication Flowchart',
        '   3.3 Farmer Verification Flowchart',
        '   3.4 Order Processing Flowchart',
        '   3.5 Data Flow Diagram (Level-0)',
        '   3.6 Data Flow Diagram (Level-1)',
        '4. Database Design',
        '5. Module Breakdown',
        '   5.1 Customer Portal',
        '   5.2 Farmer Dashboard',
        '   5.3 Administrator Panel',
        '6. Test Data Used in the Project',
        '7. Project Installation Instructions',
        '8. User Credentials',
    ]
    for item in toc_items:
        p = doc.add_paragraph(item)
        for run in p.runs:
            run.font.size = Pt(12)

    doc.add_page_break()

    # ==================== 1. PROBLEM DEFINITION ====================
    styled_heading(doc, '1. Problem Definition', 1)
    doc.add_paragraph("Agriculture is the backbone of many developing economies, yet small-scale farmers consistently face severe challenges in reaching end consumers. The traditional supply chain is riddled with intermediaries — wholesalers, distributors, and retailers — each adding their own margin. This results in two critical problems:")
    bullet(doc, "Farmers receive only a fraction of the final retail price for their produce, discouraging sustainable farming practices.")
    bullet(doc, "Consumers pay inflated prices for produce that has often traveled through multiple hands, losing freshness along the way.")
    doc.add_paragraph("Furthermore, small farmers lack the digital literacy and resources to establish an online presence on existing platforms like Daraz or Foodpanda, which are designed for large-scale commercial operations.")
    styled_heading(doc, 'Proposed Solution', 2)
    doc.add_paragraph("HarvestHub is a cross-platform mobile application that directly connects local farmers with health-conscious consumers. It eliminates the middleman entirely by providing a dedicated, easy-to-use digital marketplace where farmers can list their fresh produce (fruits, vegetables, dairy, grains, organic herbs) and consumers can browse, compare, and place orders directly.")
    doc.add_paragraph("The platform enforces a strict three-portal architecture — Customer, Farmer, and Administrator — each with isolated responsibilities, ensuring data security, operational clarity, and a premium user experience.")

    doc.add_page_break()

    # ==================== 2. DESIGN SPECIFICATIONS ====================
    styled_heading(doc, '2. Design Specifications', 1)
    
    styled_heading(doc, '2.1 Technology Stack', 2)
    styled_table(doc, ['Component', 'Technology', 'Purpose'], [
        ['Frontend', 'Flutter SDK (Dart)', 'Cross-platform mobile UI (Android & iOS)'],
        ['Backend', 'Firebase Suite', 'Serverless authentication, database, and storage'],
        ['Database', 'Cloud Firestore (NoSQL)', 'Real-time data sync across all devices'],
        ['File Storage', 'Firebase Storage', 'Product images, farmer verification documents'],
        ['Authentication', 'Firebase Auth', 'Secure Email/Password session management'],
        ['State Management', 'Provider', 'Reactive UI updates across the widget tree'],
        ['Routing', 'GoRouter', 'Deep-linkable, role-based secure navigation'],
        ['AI Integration', 'Google Generative AI', 'Harvi - AI Farm Assistant chatbot'],
        ['Location Services', 'Geolocator + Flutter Map', 'Farm location tagging and map display'],
        ['Image Handling', 'Image Picker + HTTP', 'Camera/gallery capture and cloud upload'],
    ])

    styled_heading(doc, '2.2 System Architecture', 2)
    doc.add_paragraph("HarvestHub follows a Three-Tier Serverless Architecture:")
    bold_line(doc, "Presentation Layer: ", "Built with Flutter, providing a natively compiled 60fps experience. Uses the Provider pattern for reactive state management and GoRouter for role-based routing that prevents unauthorized access to secure portals.")
    bold_line(doc, "Application Logic Layer: ", "Entirely serverless, handled by Firebase. Business rules are enforced through Firestore Security Rules. Firebase Authentication manages secure session tokens for all three user roles.")
    bold_line(doc, "Data Layer: ", "Cloud Firestore serves as the primary NoSQL database with real-time listeners for instant data synchronization. Firebase Storage hosts all media assets (product images, verification documents) and delivers them through a global CDN.")


    styled_heading(doc, '2.3 Color Theme & UI/UX Guidelines', 2)
    styled_table(doc, ['Element', 'Color Code', 'Usage'], [
        ['Primary', '#1E3C2F (Deep Forest Green)', 'Headers, navigation, primary buttons, branding'],
        ['Primary Container', '#2E7D32 (Medium Green)', 'Card backgrounds, secondary elements'],
        ['Accent / CTA', '#F57C00 (Harvest Orange)', 'Call-to-action buttons, alerts, verification prompts'],
        ['Background', '#F8FAF8 (Soft Off-White)', 'Main screen backgrounds'],
        ['Dark Background', '#11231A (Rich Dark Green)', 'Premium screens (Farmer Waiting Screen)'],
        ['Surface', '#FFFFFF (White)', 'Cards, dialogs, bottom sheets'],
        ['On Surface', '#212121 (Dark)', 'Primary text color'],
        ['On Surface Variant', '#616161 (Grey)', 'Secondary/caption text'],
    ])
    doc.add_paragraph("The UI design language emphasizes Glassmorphism (translucent blurred card backgrounds), micro-animations (pulsing status icons), and smooth page transitions to create a premium, modern feel throughout the application.")

    doc.add_page_break()

    # ==================== 3. DIAGRAMS ====================
    styled_heading(doc, '3. Diagrams', 1)

    styled_heading(doc, '3.1 Application Flow Diagram', 2)
    doc.add_paragraph("The overall application flow is as follows:")
    flow_lines = [
        "App Launch",
        "    |",
        "    v",
        "Splash Screen --> Check Auth State",
        "    |                    |",
        "    v                    v",
        "Not Logged In        Logged In",
        "    |                    |",
        "    v                    v",
        "Login Screen       Check User Role",
        "    |              /     |      \\",
        "    v             v      v       v",
        "Register      Customer  Farmer  Admin",
        "              Portal   Portal   Portal",
    ]
    for line in flow_lines:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)


    styled_heading(doc, '3.2 User Authentication Flowchart', 2)
    auth_flow = [
        "START",
        "  |",
        "  v",
        "User enters Email & Password",
        "  |",
        "  v",
        "Firebase Auth validates credentials",
        "  |",
        "  +--> INVALID --> Show Error Message --> RETRY",
        "  |",
        "  +--> VALID --> Fetch user document from Firestore",
        "         |",
        "         v",
        "     Read 'role' field",
        "         |",
        "         +---> 'customer' --> Navigate to Customer Home",
        "         |",
        "         +---> 'farmer' --> Check 'isApproved' flag",
        "         |         |",
        "         |         +--> false --> Show Approval Pending Screen",
        "         |         +--> true  --> Navigate to Farmer Dashboard",
        "         |",
        "         +---> 'admin' --> Navigate to Admin Dashboard",
        "  |",
        "  v",
        "END",
    ]
    for line in auth_flow:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)


    styled_heading(doc, '3.3 Farmer Verification Flowchart', 2)
    ver_flow = [
        "START: Farmer Registration",
        "  |",
        "  v",
        "Farmer fills form (Name, Farm Name, Contact, Location)",
        "  |",
        "  v",
        "Firebase Auth creates account --> Firestore doc created",
        "  (isApproved = false, verificationDocs = [])",
        "  |",
        "  v",
        "Farmer logs in --> Sees Approval Pending Screen",
        "  |",
        "  v",
        "Screen shows 'Action Required' with pulsing orange icon",
        "  |",
        "  v",
        "Farmer selects exactly 2 documents from gallery",
        "  |",
        "  v",
        "Taps 'Submit for Verification'",
        "  |",
        "  v",
        "Documents uploaded to Firebase Storage",
        "  |",
        "  v",
        "Firestore farmer doc updated: verificationDocs = [url1, url2]",
        "  |",
        "  v",
        "UI transitions to 'Documents Under Review' (blue icon)",
        "  |",
        "  v",
        "Admin reviews docs in Admin Panel",
        "  |",
        "  +--> Approve --> isApproved = true --> Farmer gains Dashboard access",
        "  +--> Reject  --> isApproved stays false --> Farmer remains pending",
        "  |",
        "  v",
        "END",
    ]
    for line in ver_flow:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)


    styled_heading(doc, '3.4 Order Processing Flowchart', 2)
    order_flow = [
        "START: Customer places order",
        "  |",
        "  v",
        "Order document created in Firestore",
        "  (status = 'Pending')",
        "  |",
        "  v",
        "Farmer receives notification on Dashboard",
        "  |",
        "  v",
        "Farmer reviews order items",
        "  |",
        "  +--> Confirm --> status = 'Confirmed'",
        "  |",
        "  v",
        "Farmer prepares order",
        "  |",
        "  +--> Ready --> status = 'Ready for Pickup'",
        "  |",
        "  v",
        "Customer picks up order",
        "  |",
        "  +--> Complete --> status = 'Completed'",
        "  |",
        "  v",
        "Customer can leave a review & rating",
        "  |",
        "  v",
        "END",
    ]
    for line in order_flow:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(10)

    doc.add_page_break()

    styled_heading(doc, '3.5 Data Flow Diagram (Level-0 / Context Diagram)', 2)
    dfd0 = [
        "+------------------+                          +------------------+",
        "|    Customer      | ----> Browse/Order -----> |                  |",
        "|   (External)     | <---- Notifications <---- |                  |",
        "+------------------+                          |                  |",
        "                                              |   HarvestHub     |",
        "+------------------+                          |    System        |",
        "|    Farmer        | ----> List Products ----> |   (Process 0)   |",
        "|   (External)     | <---- Order Updates <---- |                  |",
        "+------------------+                          |                  |",
        "                                              |                  |",
        "+------------------+                          |                  |",
        "|  Administrator   | ----> Approve/Deny -----> |                  |",
        "|   (External)     | <---- Reports <---------- |                  |",
        "+------------------+                          +------------------+",
        "                                                      |",
        "                                                      v",
        "                                              +------------------+",
        "                                              | Firebase Cloud   |",
        "                                              | Firestore (DB)   |",
        "                                              +------------------+",
    ]
    for line in dfd0:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(8)


    styled_heading(doc, '3.6 Data Flow Diagram (Level-1)', 2)
    dfd1 = [
        "Customer -----> [1.0 Authentication] -----> Firebase Auth",
        "                       |",
        "                       v",
        "Customer -----> [2.0 Browse Products] <----> Firestore (products)",
        "                       |",
        "                       v",
        "Customer -----> [3.0 Manage Cart] --------> Local State (Provider)",
        "                       |",
        "                       v",
        "Customer -----> [4.0 Place Order] ---------> Firestore (orders)",
        "                       |",
        "                       v",
        "Farmer   <----- [5.0 Fulfill Order] <------> Firestore (orders)",
        "                       |",
        "                       v",
        "Admin    -----> [6.0 Manage Platform] <----> Firestore (farmers, deals)",
        "                       |",
        "                       v",
        "Admin    -----> [7.0 Verify Farmers] <-----> Firebase Storage (docs)",
    ]
    for line in dfd1:
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)

    doc.add_page_break()

    # ==================== 4. DATABASE DESIGN ====================
    styled_heading(doc, '4. Database Design', 1)
    doc.add_paragraph("HarvestHub uses Cloud Firestore, a NoSQL document-oriented database. The schema is optimized for read-heavy operations with denormalized data structures to minimize query complexity.")

    styled_heading(doc, '4.1 users Collection', 2)
    styled_table(doc, ['Field', 'Type', 'Description'], [
        ['uid', 'String (PK)', 'Unique user identifier from Firebase Auth'],
        ['name', 'String', 'Full name of the user'],
        ['email', 'String', 'Email address used for authentication'],
        ['role', 'String', 'User role: "customer", "farmer", or "admin"'],
        ['createdAt', 'Timestamp', 'Account creation date and time'],
    ])

    styled_heading(doc, '4.2 customers Collection', 2)
    styled_table(doc, ['Field', 'Type', 'Description'], [
        ['uid', 'String (FK)', 'References users.uid'],
        ['name', 'String', 'Customer display name'],
        ['phone', 'String', 'Contact number'],
        ['address', 'String', 'Delivery/pickup address'],
        ['wishlist', 'Array<String>', 'List of saved product IDs'],
        ['profileImageUrl', 'String', 'URL to profile photo in Firebase Storage'],
    ])

    styled_heading(doc, '4.3 farmers Collection', 2)
    styled_table(doc, ['Field', 'Type', 'Description'], [
        ['uid', 'String (FK)', 'References users.uid'],
        ['farmName', 'String', 'Name of the farm business'],
        ['location', 'String', 'Farm address / city'],
        ['description', 'String', 'Farm description and specialties'],
        ['contactNumber', 'String', 'Business contact number'],
        ['rating', 'Number', 'Average rating from customer reviews'],
        ['isApproved', 'Boolean', 'Admin approval status (default: false)'],
        ['isSuspended', 'Boolean', 'Suspension flag (default: false)'],
        ['verificationDocs', 'Array<String>', 'Firebase Storage URLs of uploaded ID/photos'],
        ['latitude', 'Number', 'GPS latitude of farm location'],
        ['longitude', 'Number', 'GPS longitude of farm location'],
    ])

    styled_heading(doc, '4.4 products Collection', 2)
    styled_table(doc, ['Field', 'Type', 'Description'], [
        ['productId', 'String (PK)', 'Auto-generated unique product identifier'],
        ['farmerId', 'String (FK)', 'References farmers.uid'],
        ['title', 'String', 'Product display name'],
        ['category', 'String', 'Product category (Vegetables, Fruits, etc.)'],
        ['price', 'Number', 'Price per unit in PKR'],
        ['stock_qty', 'Number', 'Available quantity (0 = out of stock)'],
        ['unit', 'String', 'Unit of measurement (kg, dozen, litre)'],
        ['imageUrl', 'String', 'Product image URL from Firebase Storage'],
        ['averageRating', 'Number', 'Aggregated customer rating'],
    ])

    styled_heading(doc, '4.5 orders Collection', 2)
    styled_table(doc, ['Field', 'Type', 'Description'], [
        ['orderId', 'String (PK)', 'Auto-generated order identifier'],
        ['customerId', 'String (FK)', 'References customers.uid'],
        ['farmerId', 'String (FK)', 'References farmers.uid'],
        ['items', 'Array<Map>', 'List of ordered items with qty and price'],
        ['totalPrice', 'Number', 'Total order amount in PKR'],
        ['status', 'String', 'Pending / Confirmed / Ready / Completed'],
        ['pickupSlot', 'String', 'Selected pickup time slot'],
        ['createdAt', 'Timestamp', 'Order placement timestamp'],
    ])

    doc.add_page_break()

    # ==================== 5. MODULE BREAKDOWN ====================
    styled_heading(doc, '5. Module Breakdown', 1)

    styled_heading(doc, '5.1 Customer Portal', 2)
    doc.add_paragraph("The Customer Portal is designed to feel like a premium digital grocery store with the following features:")
    bullet(doc, "Dynamic Home Screen: Displays product categories (Vegetables, Fruits, Dairy, Organic Herbs, Grains), active deals of the day, and nearby farmer recommendations.")
    bullet(doc, "Advanced Search & Filtering: Customers can search products by name, filter by category, and sort by price or rating.")
    bullet(doc, "Product Detail View: Shows high-resolution product images, pricing, available stock, farmer information (name, location, rating), and customer reviews.")
    bullet(doc, "Shopping Cart: Full cart management with quantity adjustments, item removal, and real-time total calculation.")
    bullet(doc, "Wishlist: Customers can save products for future purchase.")
    bullet(doc, "Checkout & Pickup Slot Selection: Customers select from farmer-defined pickup time slots and place the order.")
    bullet(doc, "Order Tracking: View order status updates in real-time (Pending, Confirmed, Ready for Pickup, Completed).")
    bullet(doc, "Harvi - AI Farm Assistant: An intelligent chatbot powered by Google Generative AI that answers nutrition questions, provides storage tips, and offers seasonal farming advice.")
    bullet(doc, "Review & Rating System: After order completion, customers can rate products and leave detailed reviews.")

    styled_heading(doc, '5.2 Farmer Dashboard', 2)
    doc.add_paragraph("The Farmer Dashboard empowers agricultural sellers with enterprise-grade tools in a simple interface:")
    bullet(doc, "Secure Onboarding: Farmers register with basic details (name, farm name, contact, location). Their account is created with isApproved = false.")
    bullet(doc, "Premium Verification Screen: Upon login, unapproved farmers see a stunning glassmorphic waiting screen. A pulsing orange icon signals 'Action Required'. They must upload exactly 2 verification documents (ID Card and Farm Photo). Upon submission, the UI transitions to a blue 'Documents Under Review' state with a disabled 'Documents Submitted' button.")
    bullet(doc, "Inventory Management: Approved farmers can add new products with images, set prices, define stock quantities, and edit or delete existing listings.")
    bullet(doc, "Order Management: Farmers receive real-time order notifications and can transition orders through the lifecycle: Pending -> Confirmed -> Ready for Pickup -> Completed.")
    bullet(doc, "Sales Analytics: Dashboard overview showing total orders, revenue summaries, and product performance metrics.")

    styled_heading(doc, '5.3 Administrator Panel', 2)
    doc.add_paragraph("The Administrator Panel provides complete platform oversight with streamlined controls:")
    bullet(doc, "Farmer Verification Workflow: Admins view pending farmer applications in a 'Pending Review' tab. Tapping any farmer card opens a detailed view showing their business information and uploaded verification documents. Admins can expand documents to full-screen for review and approve or suspend accounts with a single tap.")
    bullet(doc, "Farmer Directory Management: All verified farmers are listed with their status (Active / Suspended). Admins can suspend, reactivate, or permanently delete farmer accounts.")
    bullet(doc, "Deals of the Day Management: The manage module is focused exclusively on the 'Deals of the Day' feature, allowing admins to highlight specific products on the customer homepage to drive targeted sales and engagement.")

    doc.add_page_break()

    # ==================== 6. TEST DATA ====================
    styled_heading(doc, '6. Test Data Used in the Project', 1)
    doc.add_paragraph("The application includes seed scripts to populate the database with realistic test data for demonstration and evaluation purposes. The following test data is pre-loaded:")

    styled_heading(doc, '6.1 Test User Accounts', 2)
    styled_table(doc, ['Role', 'Email', 'Password', 'Purpose'], [
        ['Customer', 'customer@example.com', 'password123', 'Browse products, place orders, use AI assistant'],
        ['Farmer', 'farmer@example.com', 'password123', 'Manage inventory, fulfill orders'],
        ['Admin', 'admin@example.com', 'password123', 'Approve farmers, manage deals'],
    ])

    styled_heading(doc, '6.2 Test Product Categories', 2)
    styled_table(doc, ['Category', 'Sample Products', 'Price Range (PKR)'], [
        ['Vegetables', 'Organic Tomatoes, Fresh Spinach, Red Onions', '80 - 250'],
        ['Fruits', 'Sindhi Mangoes, Balochistan Apples, Citrus Kinnow', '150 - 500'],
        ['Dairy & Eggs', 'Fresh Farm Milk, Desi Eggs, Buffalo Yogurt', '120 - 350'],
        ['Organic Herbs', 'Fresh Mint, Coriander, Green Chilies', '40 - 150'],
        ['Grains & Pulses', 'Basmati Rice, Red Lentils, Chickpeas', '200 - 600'],
        ['Farm Honey', 'Pure Sidr Honey, Acacia Honey', '800 - 2500'],
        ['Cold Pressed Oils', 'Mustard Oil, Coconut Oil', '400 - 1200'],
        ['Meat & Poultry', 'Farm Fresh Chicken, Mutton', '500 - 2000'],
    ])

    styled_heading(doc, '6.3 Test Farmer Profiles', 2)
    styled_table(doc, ['Farm Name', 'Location', 'Specialty', 'Rating'], [
        ['Khan Organic Farms', 'Malir, Karachi', 'Organic Vegetables', '4.8'],
        ['Sindh Fresh Harvest', 'Hyderabad', 'Seasonal Fruits', '4.6'],
        ['Punjab Grain House', 'Lahore', 'Premium Grains & Pulses', '4.7'],
        ['Balochistan Orchards', 'Quetta', 'Mountain Apples & Nuts', '4.9'],
    ])

    doc.add_page_break()

    # ==================== 7. INSTALLATION INSTRUCTIONS ====================
    styled_heading(doc, '7. Project Installation Instructions', 1)
    doc.add_paragraph("Follow these steps exactly to set up and run HarvestHub on your local machine:")

    styled_heading(doc, 'Prerequisites', 2)
    bullet(doc, "Flutter SDK (3.13.2 or higher) installed and added to PATH")
    bullet(doc, "Android Studio with Android SDK (for Android emulator)")
    bullet(doc, "Node.js (v18 or higher) installed for running seed scripts")
    bullet(doc, "Git installed for cloning the repository")
    bullet(doc, "A physical Android device or configured Android Emulator")

    styled_heading(doc, 'Step 1: Clone the Repository', 2)
    p = doc.add_paragraph("git clone https://github.com/sardarubaid-dev/Harvest_hub.git")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    p = doc.add_paragraph("cd Harvest_hub")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)

    styled_heading(doc, 'Step 2: Install Flutter Dependencies', 2)
    p = doc.add_paragraph("flutter pub get")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    doc.add_paragraph("This command downloads all required Dart packages defined in pubspec.yaml, including Firebase, Provider, GoRouter, and Google Generative AI.")

    styled_heading(doc, 'Step 3: Seed the Database with Test Data', 2)
    doc.add_paragraph("The application requires pre-populated data for demonstration. Run the following commands to seed Firebase with test users and products:")
    p = doc.add_paragraph("npm install firebase")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    p = doc.add_paragraph("node seed_auth.js")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    p = doc.add_paragraph("node seed_firebase.js")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)

    styled_heading(doc, 'Step 4: Run the Application', 2)
    doc.add_paragraph("Connect a physical device via USB or start an Android emulator, then run:")
    p = doc.add_paragraph("flutter run")
    for run in p.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    doc.add_paragraph("The application will compile and launch on the connected device. The first build may take 2-3 minutes.")

    doc.add_page_break()

    # ==================== 8. USER CREDENTIALS ====================
    styled_heading(doc, '8. User Credentials (Login ID & Password)', 1)
    doc.add_paragraph("The following credentials are pre-configured in the system for testing and evaluation purposes. Use these to access each portal:")

    styled_table(doc, ['Portal', 'Login Email', 'Password', 'Access Level'], [
        ['Customer App', 'customer@example.com', 'password123', 'Full buyer access: browse, cart, orders, AI chat'],
        ['Farmer Dashboard', 'farmer@example.com', 'password123', 'Full seller access: inventory, orders, analytics'],
        ['Admin Panel', 'admin@example.com', 'password123', 'Full admin access: approvals, deals, management'],
    ])

    doc.add_paragraph("Note: New users can also register through the app's registration screen. New farmer accounts will be placed in 'Pending Approval' status until an administrator reviews and approves their verification documents.")

    # ==================== SAVE ====================
    output_path = '/home/muhammad-hanzala/Downloads/techwiz/Harvest_hub/documentation(Harvest_HUB).docx'
    doc.save(output_path)
    print(f"Document successfully created: {output_path}")

if __name__ == "__main__":
    create_doc()
