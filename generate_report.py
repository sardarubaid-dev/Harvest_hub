import sys
import subprocess
try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--break-system-packages"])
    import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_premium_docx(template_path):
    try:
        doc = docx.Document(template_path)
    except Exception:
        doc = docx.Document()
    
    # Add HarvestHub Header
    header = doc.add_heading('HarvestHub: Premium Project Documentation', 0)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph("A state-of-the-art, multi-platform e-commerce application designed to bridge the gap between local farmers and everyday consumers.")
    
    # Add Tech Stack
    doc.add_heading('Technology Stack', level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Component'
    hdr_cells[1].text = 'Technology'
    
    stack = [
        ('Frontend', 'Flutter (Cross-platform)'),
        ('Backend', 'Firebase (Firestore, Auth, Storage)'),
        ('State Management', 'Provider'),
        ('Routing', 'GoRouter'),
        ('Mapping', 'Flutter Map & Geolocator')
    ]
    for comp, tech in stack:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = tech

    doc.add_heading('System Architecture', level=1)
    doc.add_paragraph("The architecture of HarvestHub follows a standard Client-Server model, heavily utilizing Firebase's Serverless infrastructure to handle real-time data syncing, authentication, and secure file storage.")
    
    doc.add_heading('Database Schema', level=1)
    doc.add_paragraph("Optimized NoSQL Firestore Database structure.")
    
    doc.add_heading('Users', level=2)
    doc.add_paragraph("uid, name, email, phone, role (customer/admin)")
    
    doc.add_heading('Farmers', level=2)
    doc.add_paragraph("uid, farmName, location, description, rating")
    
    doc.add_heading('Products', level=2)
    doc.add_paragraph("productId, farmerId, name, price, quantity, averageRating")
    
    doc.add_heading('Orders', level=2)
    doc.add_paragraph("orderId, customerId, farmerId, items, status, totalAmount")
    
    doc.add_heading('Reviews', level=2)
    doc.add_paragraph("reviewId, productId, userId, rating, comment, mediaUrls")
    
    doc.add_heading('UI/UX Design System', level=1)
    p = doc.add_paragraph()
    p.add_run("Primary Color: ").bold = True
    p.add_run("Deep Forest Green (#2E7D32)\n")
    p.add_run("Background: ").bold = True
    p.add_run("Soft Off-White (#F9FBF9)\n")
    p.add_run("Accents: ").bold = True
    p.add_run("Harvest Orange (#F57C00)")

    doc.add_heading('Project Credentials & Configuration', level=1)
    doc.add_paragraph("The following are the actual API keys and configuration settings used in the HarvestHub production environment. Keep this information secure.")
    
    doc.add_heading('Firebase Configuration', level=2)
    fb_table = doc.add_table(rows=1, cols=2)
    fb_hdr = fb_table.rows[0].cells
    fb_hdr[0].text = 'Setting'
    fb_hdr[1].text = 'Value'
    
    fb_config = [
        ('Project ID', 'harvest-hub-9ccfb'),
        ('App ID (Web)', '1:970168239081:web:4dea5acf961ea6a3524178'),
        ('App ID (Android)', '1:970168239081:android:9b6079d15c0467bb524178'),
        ('API Key', 'AIzaSyAxFgsMAXG6Uq8e6V7VRQotrG4OIPz93Ng'),
        ('Auth Domain', 'harvest-hub-9ccfb.firebaseapp.com'),
        ('Storage Bucket', 'harvest-hub-9ccfb.firebasestorage.app')
    ]
    for k, v in fb_config:
        row_cells = fb_table.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = v

    doc.add_heading('Conclusion', level=1)
    doc.add_paragraph("HarvestHub successfully demonstrates a robust, scalable solution for modern agricultural e-commerce.")

    doc.save('HarvestHub_Premium_Report.docx')
    print("Successfully created HarvestHub_Premium_Report.docx")

create_premium_docx('Project Report Document.docx')
